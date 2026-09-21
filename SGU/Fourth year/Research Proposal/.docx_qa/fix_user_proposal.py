from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = Path(r"C:\Users\Admin\Documents\Obsidian\SGU\Fourth year\Research Proposal\Lâm Thái Yến Nhi_Đềcương_ĐACN (1).docx")
OUT = Path(r"C:\Users\Admin\Documents\Obsidian\SGU\Fourth year\Research Proposal\.docx_qa\proposal_corrected.docx")


def set_text(p, text):
    for r in list(p.runs):
        p._p.remove(r._r)
    p.add_run(text)


def insert_before(anchor, text, style):
    marker = OxmlElement("w:p")
    anchor._p.addprevious(marker)
    p = anchor._parent.add_paragraph()
    p._p.getparent().remove(p._p)
    marker.getparent().replace(marker, p._p)
    p.style = style
    p.add_run(text)
    return p


doc = Document(SRC)

# Remove the two empty body paragraphs that forced each full-page cover table
# onto the following page, producing two unintended blank pages.
body = doc._element.body
original_children = list(body)
for child in [original_children[0], original_children[2]]:
    if child.tag == qn("w:p") and not "".join(child.itertext()).strip():
        body.remove(child)

paras = doc.paragraphs

replacements = {
    "Ràng buộc chính gồm tính duy nhất của phép gán, giới hạn sức chứa và giới hạn thời gian đến điểm đón. Vì mô hình không ràng buộc sức chứa; đây là hướng mở rộng sau nghiên cứu.":
        "Ràng buộc chính gồm tính duy nhất của phép gán, giới hạn sức chứa và giới hạn thời gian đến điểm đón. Vì mô hình không xét đi chung xe hoặc một tài xế phục vụ chuỗi nhiều yêu cầu trong cùng Batch, các ràng buộc tải trọng biến đổi và sắp xếp nhiều điểm đón - trả không được đưa vào mô hình cơ sở; đây là hướng mở rộng sau nghiên cứu.",
    "Mục tiêu này nhằm tối thiểu hóa tổng quãng đường chạy rỗng của tất cả các tài xế được hệ thống lựa chọn để phân công trong Batch hiện tại.Quãng đường chạy rỗng được tính từ vị trí hiện tại của tài xế tại thời điểm chốt Batch cho đến điểm đón của hành khách được gán cho tài xế đó.Việc tối ưu hàm này giúp hệ thống tiết kiệm nhiên liệu, giảm chi phí hao mòn cho tài xế và rút ngắn thời gian di chuyển không sinh lời.":
        "Mục tiêu này nhằm tối thiểu hóa tổng quãng đường chạy rỗng của tất cả các tài xế được hệ thống lựa chọn để phân công trong Batch hiện tại. Quãng đường chạy rỗng được tính từ vị trí hiện tại của tài xế tại thời điểm chốt Batch đến điểm đón của hành khách được gán. Việc tối ưu hàm này góp phần giảm thời gian và chi phí di chuyển không chở khách.",
    "Mục tiêu này nhằm tối thiểu hóa mức độ chênh lệch (độ lệch chuẩn) về tỷ lệ tiếp cận cuốc xe giữa nhóm tài xế có hoạt động trong hệ thống.Chỉ số này chỉ tính toán dựa trên những tài xế thực sự có mặt trực tuyến và có ít nhất một cơ hội nhận cuốc khả thi (tức là nằm trong phạm vi khoảng cách và thời gian hợp lệ) xuất hiện trong Batch.Việc tối thiểu hóa mức độ chênh lệch này giúp san đều cơ hội nhận chuyến, tránh tình trạng thuật toán tập trung dồn cuốc cho một vài tài xế ở vị trí quá thuận lợi và bỏ rơi các tài xế khác, từ đó bảo vệ quyền lợi thu nhập đồng đều hơn.":
        "Mục tiêu này nhằm tối thiểu hóa độ lệch chuẩn của tỷ lệ tiếp cận cuốc giữa các tài xế. Chỉ số được tính cho những tài xế đang trực tuyến và có ít nhất một yêu cầu khả thi trong Batch. Việc giảm độ phân tán của tỷ lệ này hướng tới phân bổ cơ hội nhận cuốc cân bằng hơn; thu nhập được theo dõi riêng như một chỉ số bổ trợ, không được đồng nhất với công bằng về cơ hội.",
    "Hệ thống cần tìm kiếm các phương án xếp xe sao cho đồng thời tối thiểu hóa cả hai mục tiêu nêu trên.Do hai mục tiêu này thường xuyên xung đột với nhau (ví dụ: việc cố gắng san đều cuốc xe cho tài xế ở xa có thể làm tăng tổng quãng đường chạy rỗng của toàn hệ thống), thuật toán sẽ không tìm ra một nghiệm duy nhất hoàn hảo. Thay vào đó, nó sẽ trả về một tập hợp các phương án đánh đổi tối ưu để người điều hành cân nhắc lựa chọn tùy theo chiến lược kinh doanh.":
        "Hệ thống tìm kiếm các phương án đồng thời giảm hai mục tiêu nêu trên. Do hai mục tiêu có thể xung đột, bài toán không giả định tồn tại một nghiệm duy nhất tốt nhất cho mọi tiêu chí. Thay vào đó, thuật toán tạo tập nghiệm không trội để thể hiện các mức đánh đổi, từ đó hỗ trợ lựa chọn phương án điều phối phù hợp.",
    "Để đảm bảo không xảy ra tranh chấp hoặc trùng lặp, mỗi yêu cầu đặt xe của hành khách trong một Batch chỉ được phép phân công cho tối đa một tài xế.Một yêu cầu cũng có thể không được gán cho ai nếu hệ thống không tìm thấy tài xế nào phù hợp hoặc việc gán xe vi phạm các ràng buộc khác.":
        "Để tránh gán trùng, mỗi yêu cầu đặt xe trong một Batch chỉ được phân công cho tối đa một tài xế. Yêu cầu có thể chưa được gán nếu hệ thống không tìm thấy tài xế phù hợp hoặc nếu mọi phép gán đều vi phạm ràng buộc.",
    "Hệ thống chỉ chấp nhận ghép cặp giữa tài xế và hành khách nếu số lượng người đi của yêu cầu đó nhỏ hơn hoặc bằng số ghế trống khả dụng của phương tiện tương ứng.Ràng buộc này tự động loại bỏ các cặp ghép bất khả thi về mặt vật lý. Ví dụ, một yêu cầu đặt xe đi nhóm 3 người sẽ được lập trình định cấu hình bằng 0 đối với xe máy (không thể nhận), nhưng có thể được gán cho ô tô 4 chỗ hoặc 7 chỗ nếu xe đó còn đủ ghế trống.":
        "Hệ thống chỉ chấp nhận ghép cặp khi số hành khách của yêu cầu không vượt quá sức chứa khả dụng của phương tiện. Ràng buộc này loại bỏ các cặp ghép không khả thi về mặt vật lý. Ví dụ, yêu cầu của nhóm 3 người không thể được gán cho xe máy nhưng có thể được gán cho ô tô nếu xe còn đủ chỗ hành khách.",
    "Nếu thuật toán quyết định gán một tài xế cho một yêu cầu, thì thời điểm tài xế di chuyển đến điểm đón thực tế không được vượt quá giới hạn chịu đựng mặc định của hệ thống.Thời gian đến điểm đón được tính bằng thời điểm chốt Batch cộng với thời gian di chuyển dự kiến dựa trên vận tốc thực tế của phương tiện.":
        "Nếu thuật toán gán một tài xế cho một yêu cầu, thời điểm tài xế dự kiến đến điểm đón không được vượt quá giới hạn chờ tối đa. Thời điểm đến được tính từ thời điểm chốt Batch cộng với thời gian di chuyển dự kiến đến điểm đón.",
    "Đối với các yêu cầu đặt xe không thể tìm được tài xế phù hợp trong Batch hiện tại (do thiếu xe, vi phạm sức chứa hoặc vi phạm thời gian chờ), hệ thống sẽ giữ lại và chuyển các yêu cầu này sang Batch kế tiếp để tiếp tục tìm kiếm ở phiên sau, và ghi nhận là hủy chuyến nếu đã quá hạn.Để đảm bảo tính khách quan và khoa học khi đánh giá, quy tắc xử lý này bắt buộc phải được áp dụng hoàn toàn đồng nhất giữa giải thuật chính NSGA-II và phương pháp đối chứng Greedy.":
        "Yêu cầu chưa tìm được tài xế phù hợp trong Batch hiện tại được giữ lại cho Batch kế tiếp nếu vẫn còn trong thời hạn chờ; yêu cầu quá hạn được ghi nhận là không được phục vụ. Quy tắc này được áp dụng giống nhau cho NSGA-II và Greedy để bảo đảm điều kiện so sánh nhất quán.",
    "1.1. Giới thiệu chi tiết bài toán điều phối gọi xe động và các thuộc tính nền tảng (1.5 điểm)":
        "1.1. Giới thiệu chi tiết bài toán điều phối gọi xe động và các thuộc tính nền tảng",
    "1.1.4. Phân tích các ràng buộc vật lý liên quan: Ràng buộc thứ tự đón trước – trả sau, ràng buộc cửa sổ thời gian chờ tối đa của khách, và giới hạn sức chứa ghế trống của phương tiện.":
        "1.1.4. Phân tích quy trình phục vụ từ điểm đón đến điểm trả và các ràng buộc của mô hình: tính duy nhất của phép gán, cửa sổ thời gian đón khách và sức chứa phương tiện.",
}

for p in doc.paragraphs:
    if p.text.strip() in replacements:
        set_text(p, replacements[p.text.strip()])

# Restore the missing method subsection between 5.3 and 5.5.
anchor = next(p for p in doc.paragraphs if p.text.strip().startswith("5.5. Thực nghiệm"))
insert_before(anchor, "5.4. Thiết kế thuật toán NSGA-II", "Heading 2")
insert_before(anchor, "5.4.1. Mã hóa nghiệm và khởi tạo quần thể", "Heading 3")
insert_before(anchor, "Mỗi nhiễm sắc thể biểu diễn phương án gán yêu cầu cho tài xế bằng mã hóa số nguyên. Quần thể ban đầu kết hợp nghiệm ngẫu nhiên khả thi với nghiệm được gieo từ Heuristic khoảng cách nhằm cải thiện chất lượng ban đầu mà vẫn duy trì tính đa dạng.", "Normal")
insert_before(anchor, "5.4.2. Xếp hạng, lựa chọn và toán tử tiến hóa", "Heading 3")
insert_before(anchor, "Các nghiệm được xếp hạng bằng sắp xếp không trội và khoảng cách mật độ. Lựa chọn giải đấu nhị phân, lai ghép và đột biến được áp dụng trên mã hóa số nguyên để tạo quần thể con.", "Normal")
insert_before(anchor, "5.4.3. Cơ chế kiểm tra và sửa nghiệm", "Heading 3")
insert_before(anchor, "Sau mỗi toán tử, cơ chế sửa nghiệm loại bỏ gán trùng, loại bỏ cặp ghép vi phạm sức chứa hoặc giới hạn thời gian đến điểm đón, rồi thử tái phân công các yêu cầu còn trống bằng tài xế còn rảnh và khả thi.", "Normal")

# Use the correct heading hierarchy for top-level sections 6 and 10.
for p in doc.paragraphs:
    if p.text.strip() == "6. GIẢ THUYẾT KHOA HỌC":
        p.style = "Heading 1"
    if p.text.strip() == "10. DANH MỤC TÀI LIỆU THAM KHẢO":
        p.style = "Heading 1"

# Replace over-assertive, weakly supported hypothesis rationale with testable wording.
hyp = {
    "Cơ sở thực tế: Các nghiên cứu lớn về nền tảng gọi xe (như hệ thống của DiDi Chuxing hay thuật toán 2FairGA của dòng nghiên cứu KDD) đã chứng minh: việc gom yêu cầu lại để giải quyết đồng thời luôn tạo ra không gian tối ưu lớn hơn việc gán xe lập tức (Greedy). [6] [7]":
        "Cơ sở khoa học: Cơ chế Batching cho phép xem xét đồng thời nhiều cặp tài xế - yêu cầu trong một cửa sổ quyết định, trong khi các nghiên cứu về điều phối gọi xe cũng cho thấy hiệu quả và công bằng cần được đánh giá trong cùng một khung ra quyết định [4], [6]. Tuy nhiên, lợi ích cụ thể phụ thuộc vào dữ liệu, độ dài Batch và cấu hình thuật toán.",
    "Giả thuyết H₁: Khi chạy mô phỏng trên cùng một tập dữ liệu chuyến đi thực tế, phương pháp gom cụm thời gian ngắn (Batching) tích hợp thuật toán tìm kiếm Pareto (NSGA-II) được dự đoán sẽ tìm ra các phương án điều phối vượt trội hơn hẳn chiến lược tham lam (Greedy). Sự vượt trội này thể hiện ở chỗ: hệ thống có thể đồng thời giảm được tổng quãng đường xe chạy rỗng đi đón khách và thu hẹp được mức chênh lệch cơ hội nhận chuyến giữa các tài xế, mà không làm suy giảm tỷ lệ phục vụ hành khách.[8]":
        "Giả thuyết H₁: Trên cùng dữ liệu và điều kiện mô phỏng, Batching kết hợp NSGA-II dự kiến tạo được ít nhất một phương án không bị phương án Greedy chi phối xét theo hai mục tiêu quãng đường chạy rỗng và độ lệch tỷ lệ tiếp cận cuốc, đồng thời duy trì tỷ lệ phục vụ ở mức chấp nhận được.",
    "Cơ sở thực tế: Trong lý thuyết tối ưu đô thị, lợi ích kinh tế của hệ thống (quãng đường chạy rỗng ngắn nhất) và tính an sinh xã hội (chia đều cuốc cho tài xế ở xa) là hai mục tiêu mâu thuẫn trực tiếp (xung đột lợi ích giữa Rider và Driver). [9]":
        "Cơ sở khoa học: Hiệu quả vận hành và công bằng trong hệ thống gọi xe có thể xung đột vì ưu tiên tài xế gần điểm đón giúp giảm quãng đường chạy rỗng nhưng có thể làm cơ hội nhận cuốc tập trung vào một số tài xế [5] - [7].",
    "Giả thuyết H₂: Thuật toán NSGA-II được dự đoán sẽ phân tách thành công sự xung đột này và trả về một tập nghiệm không trội có hình dáng biên dạng đường cong Pareto phân bố đều. Tập nghiệm này chứng minh một quy luật thực tế: Nếu người điều hành muốn tăng tính công bằng cho tài xế lên một mức nhất định, hệ thống bắt buộc phải chấp nhận một tỷ lệ gia tăng tương ứng về tổng quãng đường chạy rỗng của đội xe.":
        "Giả thuyết H₂: NSGA-II dự kiến tạo được tập nghiệm không trội có độ phân tán đủ để quan sát mức đánh đổi giữa quãng đường chạy rỗng và công bằng theo cơ hội nhận cuốc. Hình dạng và mức đánh đổi của biên Pareto sẽ được xác định từ kết quả thực nghiệm, không giả định trước là tuyến tính hoặc luôn phân bố đều.",
    "Cơ sở thực tế: Lớp bài toán định tuyến có ràng buộc thời gian cứng (Time Windows) có không gian nghiệm cực kỳ nhạy cảm. Các phép toán cắt ghép gen ngẫu nhiên của Giải thuật di truyền (GA) thông thường có tỷ lệ làm hỏng nghiệm (vi phạm giờ đón) lên đến hơn 80%, khiến thuật toán bị sa lầy. [10]":
        "Cơ sở khoa học: Các toán tử lai ghép và đột biến trên mã hóa số nguyên có thể tạo ra nghiệm gán trùng hoặc vi phạm ràng buộc sức chứa và thời gian. Vì vậy, cần kiểm tra vai trò của cơ chế sửa nghiệm thay vì giả định mọi cá thể sinh ra đều khả thi.",
    "Giả thuyết H₃: Cơ chế sửa nghiệm chủ động (loại bỏ gán trùng và tái phân công tham lam vào các vị trí trống) được dự đoán sẽ giúp duy trì tỷ lệ các phương án hợp lệ (khả thi) trong quần thể ở mức cao qua các thế hệ. Kết quả thực nghiệm dự kiến sẽ chứng minh cơ chế này giúp thuật toán hội tụ về biên Pareto nhanh hơn và cho ra nghiệm chất lượng hơn hẳn so với việc chỉ sử dụng hàm phạt toán học để loại bỏ các nghiệm lỗi.":
        "Giả thuyết H₃: Cơ chế sửa nghiệm chủ động dự kiến làm tăng tỷ lệ cá thể khả thi so với cấu hình chỉ dùng hàm phạt, đồng thời không làm suy giảm đáng kể độ đa dạng của tập nghiệm. Giả thuyết được kiểm tra bằng thí nghiệm loại bỏ thành phần.",
    "Cơ sở thực tế: Các thuật toán Metaheuristic như NSGA-II vốn chạy chậm khi dữ liệu lớn. Tuy nhiên, các nghiên cứu chia cắt thời gian (Rolling Horizon Framework) chỉ ra rằng khi chặt nhỏ dữ liệu theo từng Batch từ 1-2 phút, quy mô thực thể trong mỗi phiên xử lý là rất nhỏ.":
        "Cơ sở khoa học: Khả năng áp dụng theo thời gian thực phụ thuộc vào quy mô mỗi Batch, kích thước quần thể, số thế hệ và chi phí đánh giá nghiệm. Do đó, thời gian tính toán phải được đo trên nhiều lần chạy và báo cáo theo phân phối.",
    "Giả thuyết H₄: Với quy mô phân luồng dữ liệu của đề tài, thời gian xử lý tính toán của thuật toán NSGA-II tại các tình huống giao dịch cao điểm nhất (được đo ở mức phân vị số 95) được dự đoán sẽ hoàn thành trong vòng dưới 60 giây. Điều này chứng minh thuật toán hoàn toàn có khả năng nhúng vào các ứng dụng gọi xe trực tuyến chạy theo thời gian thực mà không gây ra hiện tượng nghẽn hay trễ hệ thống.":
        "Giả thuyết H₄: Với quy mô Batch và cấu hình được lựa chọn trong phạm vi đề tài, phân vị 95 của thời gian tính toán dự kiến nhỏ hơn độ dài Batch 60 giây. Kết quả chỉ cho phép kết luận trong môi trường phần cứng, dữ liệu và tham số đã thử nghiệm.",
    "Cơ sở thực tế: Đây là quy luật \"Đánh đổi thời gian\" kinh điển trong Logistics. Cửa sổ thời gian tích lũy càng lâu thì thuật toán càng có nhiều dữ liệu để tối ưu tổng thể, nhưng khách hàng phải đợi phản hồi lâu hơn.":
        "Cơ sở khoa học: Batch dài hơn có thể cung cấp nhiều lựa chọn ghép nối hơn nhưng cũng làm tăng độ trễ trước khi ra quyết định; tác động ròng cần được đo bằng thực nghiệm trên cùng luồng dữ liệu [4].",
    "Giả thuyết H₅: Việc tăng độ dài cửa sổ Batch từ 60 giây lên 120 giây được dự đoán sẽ làm tăng số lượng hành khách và tài xế trong một phiên tối ưu. Thực nghiệm sẽ chứng minh điều này giúp NSGA-II tìm ra tập nghiệm Pareto có chất lượng tốt hơn (xe chạy rỗng ít hơn, công bằng hơn), nhưng sẽ làm tăng thời gian xử lý của máy tính và kéo dài tổng thời gian chờ đợi phản hồi của hành khách.":
        "Giả thuyết H₅: Khi tăng độ dài Batch từ 60 giây lên 120 giây, số cặp ghép được xem xét dự kiến tăng, có thể cải thiện một số chỉ số của tập nghiệm Pareto nhưng đồng thời làm tăng độ trễ phản hồi và chi phí tính toán. Mức thay đổi được đánh giá bằng phân tích độ nhạy.",
}
for p in doc.paragraphs:
    if p.text.strip() in hyp:
        set_text(p, hyp[p.text.strip()])

# Replace incomplete placeholder references with two complete, relevant IEEE entries.
refs = [p for p in doc.paragraphs if p.text.strip().startswith("[")]
for p in list(refs):
    n = p.text.strip().split("]", 1)[0] + "]"
    if n in {"[6]", "[7]", "[8]", "[9]", "[10]"}:
        p._element.getparent().remove(p._element)
anchor_ref = next(p for p in doc.paragraphs if p.text.strip().startswith("[5]"))
insert_after = anchor_ref
def add_after(p, text):
    q = OxmlElement("w:p")
    p._p.addnext(q)
    np = p._parent.add_paragraph()
    np._p.getparent().remove(np._p)
    q.getparent().replace(q, np._p)
    np.style = "Normal"
    np.add_run(text)
    return np
insert_after = add_after(insert_after, "[6] J. Sun, H. Jin, Z. Yang, L. Su, and X. Wang, “Optimizing long-term efficiency and fairness in ride-hailing via joint order dispatching and driver repositioning,” in Proc. 28th ACM SIGKDD Conf. Knowledge Discovery and Data Mining, 2022, pp. 3950-3960, doi: 10.1145/3534678.3539060.")
add_after(insert_after, "[7] A. A. Makhdomi and I. A. Gillani, “Towards a greener and fairer transportation system: A survey of route recommendation techniques,” ACM Trans. Intell. Syst. Technol., vol. 15, no. 1, pp. 1:1-1:57, 2024, doi: 10.1145/3627825.")

# Add the required chapter summary missing from the projected Chapter 3 outline.
conclusion = next(p for p in doc.paragraphs if p.text.strip() == "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN")
insert_before(conclusion, "3.4. Tóm tắt Chương 3", "Normal")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
