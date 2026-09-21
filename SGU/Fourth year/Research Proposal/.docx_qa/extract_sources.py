from pathlib import Path
import json
from docx import Document
from docx.oxml.ns import qn


def paragraph_record(p, index):
    pf = p.paragraph_format
    runs = []
    for r in p.runs:
        rpr = r._element.rPr
        fonts = {}
        if rpr is not None and rpr.rFonts is not None:
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                value = rpr.rFonts.get(qn(f"w:{key}"))
                if value:
                    fonts[key] = value
        runs.append({
            "text": r.text,
            "bold": r.bold,
            "italic": r.italic,
            "underline": bool(r.underline) if r.underline is not None else None,
            "font_name": r.font.name,
            "font_size_pt": r.font.size.pt if r.font.size else None,
            "fonts_xml": fonts,
        })
    return {
        "index": index,
        "text": p.text,
        "style": p.style.name if p.style else None,
        "alignment": str(p.alignment),
        "left_indent_pt": pf.left_indent.pt if pf.left_indent else None,
        "right_indent_pt": pf.right_indent.pt if pf.right_indent else None,
        "first_line_indent_pt": pf.first_line_indent.pt if pf.first_line_indent else None,
        "space_before_pt": pf.space_before.pt if pf.space_before else None,
        "space_after_pt": pf.space_after.pt if pf.space_after else None,
        "line_spacing": str(pf.line_spacing),
        "keep_with_next": pf.keep_with_next,
        "page_break_before": pf.page_break_before,
        "runs": runs,
    }


def main():
    root = Path(__file__).resolve().parents[1]
    guide = root / "HD trinh bay de cuong ĐACN.docx"
    doc = Document(guide)
    data = {
        "paragraphs": [paragraph_record(p, i) for i, p in enumerate(doc.paragraphs)],
        "tables": [],
        "styles": [],
    }
    for ti, table in enumerate(doc.tables):
        data["tables"].append({
            "index": ti,
            "style": table.style.name if table.style else None,
            "rows": [[cell.text for cell in row.cells] for row in table.rows],
        })
    for s in doc.styles:
        if s.type == 1:
            data["styles"].append({
                "name": s.name,
                "base": s.base_style.name if s.base_style else None,
                "font": s.font.name,
                "size_pt": s.font.size.pt if s.font.size else None,
                "bold": s.font.bold,
                "italic": s.font.italic,
            })
    (root / ".docx_qa" / "guide_extract.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    lines = []
    for p in data["paragraphs"]:
        if p["text"].strip():
            lines.append(f"P{p['index']:03d} [{p['style']}] {p['text']}")
    for t in data["tables"]:
        lines.append(f"\nTABLE {t['index']} style={t['style']}")
        for row in t["rows"]:
            lines.append(" | ".join(x.replace("\n", " / ") for x in row))
    (root / ".docx_qa" / "guide_text.txt").write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
