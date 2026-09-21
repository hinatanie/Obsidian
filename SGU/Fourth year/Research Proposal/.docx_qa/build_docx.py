from pathlib import Path
import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
OUT = Path(os.environ.get("DOCX_OUT", str(ROOT / "De_cuong_nghien_cuu_NSGAII_Ride_Hailing.docx")))
MML2OMML = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")

MATH = {
    "eq_5_1.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><mo>min</mo><mspace width="0.25em"/><msub><mi>F</mi><mn>1</mn></msub><mo>(</mo><mi>X</mi><mo>)</mo><mo>=</mo><mrow><munderover><mo>∑</mo><mrow><mi>i</mi><mo>∈</mo><mi>D</mi></mrow><mrow/></munderover><munderover><mo>∑</mo><mrow><mi>j</mi><mo>∈</mo><mi>R</mi></mrow><mrow/></munderover><msub><mi>x</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>·</mo><mi>dist</mi><mo>(</mo><msub><mi>d</mi><mi>i</mi></msub><mo>,</mo><msub><mi>p</mi><mi>j</mi></msub><mo>)</mo></mrow></mrow></math>""",
    "eq_5_2.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi>q</mi><mi>i</mi></msub><mo>=</mo><mfrac><msub><mi>n</mi><mi>i</mi></msub><msub><mi>e</mi><mi>i</mi></msub></mfrac><mo>,</mo><mspace width="0.8em"/><mover><mi>q</mi><mo>¯</mo></mover><mo>=</mo><mfrac><mn>1</mn><mrow><mo>|</mo><msup><mi>D</mi><mo>+</mo></msup><mo>|</mo></mrow></mfrac><munderover><mo>∑</mo><mrow><mi>i</mi><mo>∈</mo><msup><mi>D</mi><mo>+</mo></msup></mrow><mrow/></munderover><msub><mi>q</mi><mi>i</mi></msub></mrow></math>""",
    "eq_5_3.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><mo>min</mo><mspace width="0.25em"/><msub><mi>F</mi><mn>2</mn></msub><mo>(</mo><mi>X</mi><mo>)</mo><mo>=</mo><msqrt><mfrac><mn>1</mn><mrow><mo>|</mo><msup><mi>D</mi><mo>+</mo></msup><mo>|</mo></mrow></mfrac><munderover><mo>∑</mo><mrow><mi>i</mi><mo>∈</mo><msup><mi>D</mi><mo>+</mo></msup></mrow><mrow/></munderover><msup><mrow><mo>(</mo><msub><mi>q</mi><mi>i</mi></msub><mo>−</mo><mover><mi>q</mi><mo>¯</mo></mover><mo>)</mo></mrow><mn>2</mn></msup></msqrt></mrow></math>""",
    "eq_5_4.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><mo>min</mo><mo>(</mo><msub><mi>F</mi><mn>1</mn></msub><mo>(</mo><mi>X</mi><mo>)</mo><mo>,</mo><msub><mi>F</mi><mn>2</mn></msub><mo>(</mo><mi>X</mi><mo>)</mo><mo>)</mo></mrow></math>""",
    "eq_5_5.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><munderover><mo>∑</mo><mrow><mi>i</mi><mo>∈</mo><mi>D</mi></mrow><mrow/></munderover><msub><mi>x</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>≤</mo><mn>1</mn><mo>,</mo><mspace width="1em"/><mo>∀</mo><mi>j</mi><mo>∈</mo><mi>R</mi></mrow></math>""",
    "eq_5_6.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><munderover><mo>∑</mo><mrow><mi>j</mi><mo>∈</mo><mi>R</mi></mrow><mrow/></munderover><msub><mi>x</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>≤</mo><mn>1</mn><mo>,</mo><mspace width="1em"/><mo>∀</mo><mi>i</mi><mo>∈</mo><mi>D</mi></mrow></math>""",
    "eq_5_7.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi>x</mi><mrow><mi>i</mi><mi>j</mi></mrow></msub><mo>=</mo><mn>1</mn><mo>⇒</mo><msubsup><mi>t</mi><mi>j</mi><mi>req</mi></msubsup><mo>+</mo><msubsup><mi>t</mi><mrow><mi>i</mi><mi>j</mi></mrow><mi>pickup</mi></msubsup><mo>≤</mo><msubsup><mi>t</mi><mi>j</mi><mi>max</mi></msubsup></mrow></math>""",
    "eq_5_8.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msup><mi>X</mi><mi>a</mi></msup><mo>≺</mo><msup><mi>X</mi><mi>b</mi></msup><mo>⇔</mo><mo>[</mo><mo>∀</mo><mi>k</mi><mo>,</mo><msub><mi>F</mi><mi>k</mi></msub><mo>(</mo><msup><mi>X</mi><mi>a</mi></msup><mo>)</mo><mo>≤</mo><msub><mi>F</mi><mi>k</mi></msub><mo>(</mo><msup><mi>X</mi><mi>b</mi></msup><mo>)</mo><mo>]</mo><mo>∧</mo><mo>[</mo><mo>∃</mo><mi>k</mi><mo>,</mo><msub><mi>F</mi><mi>k</mi></msub><mo>(</mo><msup><mi>X</mi><mi>a</mi></msup><mo>)</mo><mo>&lt;</mo><msub><mi>F</mi><mi>k</mi></msub><mo>(</mo><msup><mi>X</mi><mi>b</mi></msup><mo>)</mo><mo>]</mo></mrow></math>""",
    "eq_5_9.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><mi>CD</mi><mo>(</mo><mi>i</mi><mo>)</mo><mo>=</mo><munderover><mo>∑</mo><mrow><mi>k</mi><mo>=</mo><mn>1</mn></mrow><mn>2</mn></munderover><mfrac><mrow><msub><mi>F</mi><mi>k</mi></msub><mo>(</mo><mi>i</mi><mo>+</mo><mn>1</mn><mo>)</mo><mo>−</mo><msub><mi>F</mi><mi>k</mi></msub><mo>(</mo><mi>i</mi><mo>−</mo><mn>1</mn><mo>)</mo></mrow><mrow><msubsup><mi>F</mi><mi>k</mi><mi>max</mi></msubsup><mo>−</mo><msubsup><mi>F</mi><mi>k</mi><mi>min</mi></msubsup></mrow></mfrac></mrow></math>""",
    "eq_6_1.png": """<math xmlns="http://www.w3.org/1998/Math/MathML"><mrow><msub><mi>t</mi><mi>compute</mi></msub><mo>&lt;</mo><mi>ΔT</mi></mrow></math>""",
}

TITLE = ("TIẾP CẬN BẰNG THUẬT TOÁN NSGA-II CHO BÀI TOÁN ĐIỀU PHỐI GỌI XE "
         "ĐA MỤC TIÊU THEO CƠ CHẾ BATCHING CÓ RÀNG BUỘC CỬA SỔ THỜI GIAN")


def set_run_font(run, size=13, bold=None, italic=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = None


def format_para(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first=1.0,
                before=0, after=0, line=1.5, keep=False):
    p.alignment = align
    pf = p.paragraph_format
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.first_line_indent = Cm(first) if first is not None else None
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep
    pf.widow_control = True
    return p


def configure_section(section, body=True):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.orientation = 0
    if body:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(3.5)
        section.right_margin = Cm(2)
        section.header_distance = Cm(1.2)
    else:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.header_distance = Cm(1)


def add_page_border(section):
    sect_pr = section._sectPr
    pg_borders = OxmlElement("w:pgBorders")
    pg_borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "double")
        edge.set(qn("w:sz"), "12")
        edge.set(qn("w:space"), "24")
        edge.set(qn("w:color"), "000000")
        pg_borders.append(edge)
    sect_pr.append(pg_borders)


def remove_page_border(section):
    sect_pr = section._sectPr
    for node in list(sect_pr.findall(qn("w:pgBorders"))):
        sect_pr.remove(node)


def add_page_number(section, start=1):
    section.header.is_linked_to_previous = False
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r._r.addnext(fld)
    set_run_font(r, 13)
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), str(start))
    section._sectPr.append(pg_num)


def clear_header(section):
    section.header.is_linked_to_previous = False
    for p in section.header.paragraphs:
        p._element.clear_content()


def centered(doc, text="", size=16, bold=False, italic=False, before=0, after=0, keep=False):
    p = doc.add_paragraph()
    format_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=None,
                before=before, after=after, line=1.0, keep=keep)
    r = p.add_run(text)
    set_run_font(r, size, bold, italic)
    return p


def centered_in_cell(cell, text="", size=16, bold=False, italic=False, before=0, after=0):
    p = cell.add_paragraph()
    format_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=None,
                before=before, after=after, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size, bold, italic)
    return p


def spacer(doc, points):
    p = doc.add_paragraph()
    format_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=1.0)
    p.paragraph_format.space_after = Pt(points)
    set_run_font(p.add_run(" "), 1)


def cover(doc, inner=False):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.4)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell, top=420, start=260, bottom=260, end=260)
    row = table.rows[0]
    row.height = Cm(24.8)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    prevent_row_split(row)
    set_table_borders(table, color="000000", size="14")
    first = cell.paragraphs[0]
    format_para(first, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=1.0)
    set_run_font(first.add_run("ỦY BAN NHÂN DÂN THÀNH PHỐ HỒ CHÍ MINH"), 16, True)
    centered_in_cell(cell, "TRƯỜNG ĐẠI HỌC SÀI GÒN", 16, True)
    centered_in_cell(cell, "KHOA CÔNG NGHỆ THÔNG TIN", 16, True)
    centered_in_cell(cell, "[HỌ VÀ TÊN SINH VIÊN]", 16, True, before=48)
    centered_in_cell(cell, "(Tác giả đồ án chuyên ngành)", 16, False, True)
    p = centered_in_cell(cell, TITLE, 18, True, before=54)
    p.paragraph_format.line_spacing = 1.15
    centered_in_cell(cell, "ĐỀ CƯƠNG ĐỒ ÁN CHUYÊN NGÀNH", 16, True, before=52 if not inner else 38)
    centered_in_cell(cell, "NGÀNH: [BỔ SUNG TÊN NGÀNH]", 16)
    if inner:
        centered_in_cell(cell, "GIẢNG VIÊN HƯỚNG DẪN:", 15, True, before=28)
        centered_in_cell(cell, "[HỌ VÀ TÊN GIẢNG VIÊN]", 15, True)
        centered_in_cell(cell, "Thành phố Hồ Chí Minh, năm 2026", 16, True, before=34)
    else:
        centered_in_cell(cell, "Thành phố Hồ Chí Minh, năm 2026", 16, True, before=68)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    size = 14 if level == 1 else 13
    before = 10 if level == 1 else 6
    after = 4 if level == 1 else 2
    format_para(p, align=WD_ALIGN_PARAGRAPH.LEFT, first=None,
                before=before, after=after, line=1.5, keep=True)
    r = p.add_run(text)
    set_run_font(r, size, True, level >= 3)
    p.style = doc.styles[f"Heading {min(level, 3)}"]
    return p


def add_body(doc, text, *, first=1.0, italic=False, bold_lead=None):
    p = doc.add_paragraph()
    format_para(p, first=first)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, 13, True, False)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, 13, False, italic)
    else:
        set_run_font(p.add_run(text), 13, False, italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph()
    format_para(p, first=None)
    pf = p.paragraph_format
    pf.left_indent = Cm(1.0 + 0.7 * level)
    pf.first_line_indent = Cm(-0.5)
    r = p.add_run("- " + text)
    set_run_font(r, 13)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for c, (text, width) in enumerate(zip(headers, widths)):
        cell = hdr.cells[c]
        cell.width = Cm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        format_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=1.15)
        set_run_font(p.add_run(text), 12, True)
    for ri, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for c, (text, width) in enumerate(zip(row_data, widths)):
            cell = row.cells[c]
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri % 2 == 1:
                set_cell_shading(cell, "F5F9FC")
            p = cell.paragraphs[0]
            align = WD_ALIGN_PARAGRAPH.CENTER if c in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            format_para(p, align=align, first=None, line=1.15)
            set_run_font(p.add_run(text), 11.5)
    return table


def set_table_borders(table, color="A6A6A6", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            edge = OxmlElement(f"w:{edge_name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), size)
        edge.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "nil")
        borders.append(edge)
    tbl_pr.append(borders)


def add_equation(doc, filename, number, width_cm=10.8):
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    remove_table_borders(table)
    widths = [Cm(0.2), Cm(13.8), Cm(1.2)]
    for c, width in enumerate(widths):
        table.columns[c].width = width
        table.cell(0, c).width = width
        table.cell(0, c).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(table.cell(0, c), top=20, bottom=20, start=20, end=20)
    p = table.cell(0, 1).paragraphs[0]
    format_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, line=1.0)
    transform = etree.XSLT(etree.parse(str(MML2OMML)))
    mathml = etree.fromstring(MATH[filename].encode("utf-8"))
    omml = transform(mathml).getroot()
    p._p.append(parse_xml(etree.tostring(omml)))
    pn = table.cell(0, 2).paragraphs[0]
    format_para(pn, align=WD_ALIGN_PARAGRAPH.RIGHT, first=None, line=1.0)
    set_run_font(pn.add_run(f"({number})"), 13)
    return table


def add_reference_placeholder(doc, text):
    p = doc.add_paragraph()
    format_para(p, first=None)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    set_run_font(p.add_run(text), 13)
    return p


def build():
    doc = Document()
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        s = styles[name]
        s.font.name = "Times New Roman"
        s._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        s._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        s.font.color.rgb = None

    sec1 = doc.sections[0]
    configure_section(sec1, body=False)
    remove_page_border(sec1)
    clear_header(sec1)
    cover(doc, inner=False)

    doc.add_page_break()
    sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(sec2, body=False)
    remove_page_border(sec2)
    clear_header(sec2)
    cover(doc, inner=True)

    doc.add_page_break()
    sec3 = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(sec3, body=True)
    remove_page_border(sec3)
    add_page_number(sec3, 1)

    add_heading(doc, "1. LÝ DO CHỌN ĐỀ TÀI", 1)
    add_body(doc, "Sự phát triển của dịch vụ gọi xe công nghệ làm gia tăng nhu cầu xử lý đồng thời nhiều yêu cầu di chuyển trong môi trường đô thị. Ở mỗi thời điểm, hệ thống phải ghép nối tài xế đang sẵn sàng với hành khách mới phát sinh, đồng thời kiểm soát quãng đường tài xế di chuyển không chở khách và thời gian chờ của hành khách. Khi quy mô hệ thống tăng, quyết định gán chuyến riêng lẻ theo nguyên tắc tài xế gần nhất có thể tạo ra phương án tốt tại thời điểm hiện tại nhưng chưa chắc có lợi cho toàn bộ nhóm yêu cầu.")
    add_body(doc, "Cơ chế Batching tạo một khoảng đệm ngắn để gom các yêu cầu và tài xế khả dụng thành từng nhóm, sau đó giải bài toán ghép nối đồng thời. Cách tiếp cận này chuyển dòng yêu cầu liên tục thành chuỗi bài toán tĩnh tương đối, qua đó mở rộng khả năng xem xét quan hệ giữa nhiều tài xế và nhiều hành khách trong cùng một chu kỳ. Tuy nhiên, độ dài Batch cũng tạo ra đánh đổi: Batch dài hơn có thể cung cấp nhiều lựa chọn ghép nối hơn nhưng đồng thời làm tăng độ trễ quyết định.")
    add_body(doc, "Bài toán còn có tính đa mục tiêu. Nếu chỉ tối thiểu hóa quãng đường chạy rỗng, một số tài xế ở vị trí thuận lợi có thể liên tục được phân công, trong khi các tài xế khác ít có cơ hội nhận chuyến. Ngược lại, nếu ưu tiên tuyệt đối sự công bằng, hệ thống có thể phải chọn tài xế ở xa điểm đón hơn và làm tăng chi phí vận hành. Do đó, cần một phương pháp tạo ra nhiều phương án đánh đổi thay vì quy về một hàm mục tiêu duy nhất với trọng số cố định.")
    add_body(doc, "NSGA-II là thuật toán tiến hóa đa mục tiêu có cơ chế sắp xếp không trội và khoảng cách mật độ, phù hợp để tìm tập nghiệm Pareto biểu diễn các phương án cân bằng giữa hiệu quả vận hành và công bằng. Trên cơ sở đó, đề tài lựa chọn nghiên cứu “Tiếp cận bằng thuật toán NSGA-II cho bài toán điều phối gọi xe đa mục tiêu theo cơ chế Batching có ràng buộc cửa sổ thời gian”. Nghiên cứu dự kiến xây dựng mô hình ghép một tài xế với một yêu cầu trong mỗi Batch, thiết kế cơ chế mã hóa và sửa nghiệm phù hợp, sau đó đánh giá bằng mô phỏng trên dữ liệu hành trình thực tế.")
    add_body(doc, "Về ý nghĩa khoa học, đề tài dự kiến làm rõ cách biểu diễn đồng thời hai mục tiêu xung đột và các ràng buộc thời gian trong một mô hình ghép nối theo Batch. Về ý nghĩa thực tiễn, kết quả kỳ vọng cung cấp một khung thử nghiệm có thể dùng để đánh giá mức đánh đổi giữa quãng đường chạy rỗng, thời gian chờ, tỷ lệ phục vụ và mức chênh lệch cơ hội nhận cuốc giữa các tài xế. Thu nhập chỉ được theo dõi như một chỉ số phụ nhằm phát hiện tác dụng ngoài dự kiến. Các giá trị cải thiện cụ thể chỉ được kết luận sau khi hoàn tất thực nghiệm.")

    add_heading(doc, "2. LỊCH SỬ NGHIÊN CỨU VẤN ĐỀ VÀ TỔNG QUAN", 1)
    add_heading(doc, "2.1. Nền tảng từ bài toán định tuyến và cửa sổ thời gian", 2)
    add_body(doc, "Một trong những nền tảng sớm của nghiên cứu điều phối phương tiện là bài toán Truck Dispatching do Dantzig và Ramser công bố năm 1959. Nghiên cứu xem xét việc phân công một đội xe giao xăng từ kho trung tâm đến các trạm dịch vụ sao cho nhu cầu được đáp ứng và tổng quãng đường di chuyển được giảm thiểu [1]. Công trình này đặt nền móng cho lớp bài toán định tuyến phương tiện, trong đó quyết định chính là phân chia khách hàng cho các xe và xác định tuyến phục vụ.")
    add_body(doc, "Khi yêu cầu dịch vụ gắn với giới hạn thời gian, bài toán được mở rộng thành định tuyến và lập lịch có cửa sổ thời gian. Solomon xây dựng và đánh giá một nhóm heuristic cho các bài toán mà khách hàng có thời điểm phục vụ sớm nhất và muộn nhất, đồng thời cho thấy hiệu quả của phương pháp chèn trong nhiều môi trường thực nghiệm [2]. Ràng buộc cửa sổ thời gian là cơ sở phù hợp để mô tả giới hạn chờ của hành khách trong dịch vụ gọi xe.")
    add_body(doc, "Tuy có liên hệ với VRP và bài toán đón–trả, mô hình của đề tài không tối ưu một chuỗi nhiều điểm dừng cho mỗi xe. Trong mỗi Batch, một tài xế được ghép với tối đa một yêu cầu và một yêu cầu được gán cho tối đa một tài xế. Vì vậy, bài toán cơ sở gần với bài toán phân công hai phía có ràng buộc thời gian hơn là một mô hình VRPTW hoặc Pickup and Delivery Problem đầy đủ.")

    add_heading(doc, "2.2. Từ điều phối tĩnh đến điều phối gọi xe động", 2)
    add_body(doc, "Trong hệ thống gọi xe, yêu cầu hành khách và trạng thái tài xế xuất hiện, thay đổi liên tục theo thời gian. Quyết định hiện tại không chỉ ảnh hưởng đến khoảng cách đón khách mà còn làm thay đổi phân bố tài xế cho các yêu cầu tiếp theo. Do đó, điều phối gọi xe là bài toán động và ngẫu nhiên; hệ thống phải đưa ra quyết định nhanh khi chưa biết đầy đủ nhu cầu tương lai.")
    add_body(doc, "Cách tiếp cận tức thời thường gán yêu cầu ngay khi phát sinh, chẳng hạn chọn tài xế khả thi gần điểm đón nhất. Phương pháp này dễ cài đặt và có thời gian phản hồi ngắn, nhưng mang tính tham lam vì mỗi cặp ghép được quyết định riêng lẻ. Một tài xế ở vị trí thuận lợi có thể liên tục được chọn, trong khi ảnh hưởng của phép gán đối với các yêu cầu và tài xế còn lại chưa được xem xét đầy đủ.")

    add_heading(doc, "2.3. Điều phối theo cơ chế Batching", 2)
    add_body(doc, "Cơ chế Batching tạo các cửa sổ điều phối ngắn, trong đó các yêu cầu đang mở và tài xế khả dụng được tập hợp rồi ghép nối đồng thời. Nghiên cứu về hệ thống điều phối của DiDi mô tả quá trình này dưới dạng bài toán phân công tuyến tính giữa tập yêu cầu và tập tài xế trong mỗi cửa sổ [4]. So với quy tắc tài xế gần nhất được áp dụng tuần tự, Batching cho phép xem xét nhiều cặp ghép cùng lúc và tìm phương án tổng thể hơn trong phạm vi cửa sổ hiện tại.")
    add_body(doc, "Lợi ích của Batching đi kèm một đánh đổi. Cửa sổ dài hơn có thể tạo thêm lựa chọn ghép nối nhưng làm tăng độ trễ phản hồi; cửa sổ ngắn hơn phản hồi nhanh nhưng cung cấp ít thông tin để tối ưu đồng thời. Vì vậy, độ dài Batch cần được xem là tham số thực nghiệm thay vì một giá trị tối ưu cố định. Đề tài sử dụng 60 giây làm cấu hình mặc định và 120 giây cho phân tích độ nhạy, không xem hai giá trị này là chuẩn chung của mọi hệ thống gọi xe.")

    add_heading(doc, "2.4. Tối ưu đa mục tiêu và thuật toán NSGA-II", 2)
    add_body(doc, "Điều phối gọi xe thường phải cân bằng các mục tiêu có thể xung đột, chẳng hạn giảm quãng đường chạy rỗng, giảm thời gian chờ, duy trì tỷ lệ phục vụ và bảo đảm sự công bằng. Việc cộng các mục tiêu bằng trọng số cố định tạo ra một giá trị duy nhất nhưng phụ thuộc mạnh vào cách chọn trọng số và có thể che khuất các phương án đánh đổi. Cách tiếp cận Pareto giữ các mục tiêu độc lập và tạo ra tập nghiệm không trội để người ra quyết định quan sát trực tiếp mức đánh đổi.")
    add_body(doc, "NSGA-II do Deb và cộng sự đề xuất là một thuật toán tiến hóa đa mục tiêu sử dụng sắp xếp không trội nhanh, cơ chế lưu giữ cá thể ưu tú và khoảng cách mật độ để duy trì sự đa dạng của quần thể [3]. Thuật toán được sử dụng rộng rãi cho các bài toán đa mục tiêu có không gian nghiệm lớn. Tuy nhiên, không thể mặc định NSGA-II luôn tốt hơn phương pháp chính xác hoặc heuristic; hiệu quả của thuật toán cần được đánh giá theo chất lượng tập Pareto, tỷ lệ nghiệm khả thi và thời gian tính toán trên từng quy mô Batch.")
    add_body(doc, "Một số nghiên cứu điều phối quy mô lớn mở rộng từ tối ưu tổ hợp ngắn hạn sang các phương pháp học tăng cường nhằm xét giá trị dài hạn của vị trí tài xế [4]. Các hướng lai ghép giữa thuật toán tiến hóa, tìm kiếm cục bộ hoặc mô hình học máy là xu hướng đáng chú ý, nhưng nằm ngoài phạm vi cài đặt của đề tài hiện tại.")

    add_heading(doc, "2.5. Công bằng trong phân công cuốc xe", 2)
    add_body(doc, "Công bằng trong hệ thống gọi xe không có một định nghĩa duy nhất. Đối tượng được xem xét có thể là hành khách, tài xế hoặc khu vực địa lý; thước đo có thể dựa trên tỷ lệ phục vụ, thời gian chờ, thu nhập hoặc khả năng tiếp cận yêu cầu. Chẳng hạn, Ben-Gal và Tzur nghiên cứu sự đánh đổi giữa hiệu quả và công bằng địa lý đối với hành khách trong bài toán phân công gọi xe trực tuyến [5]. Điều này cho thấy mỗi nghiên cứu phải công bố rõ đối tượng và cách đo công bằng thay vì chỉ sử dụng khái niệm fairness chung chung.")
    add_body(doc, "Đề tài tập trung vào công bằng đối với tài xế theo cơ hội nhận cuốc. Với mỗi tài xế, cơ hội được xác định khi tài xế đang sẵn sàng và có ít nhất một yêu cầu khả thi trong Batch; tỷ lệ tiếp cận cuốc bằng số cuốc thực tế được phân chia cho số lần xuất hiện cơ hội đó. Cách chuẩn hóa này tránh đánh giá bất lợi cho tài xế ít thời gian hoạt động hoặc ở thời điểm không có yêu cầu khả thi. Thu nhập được theo dõi như chỉ số phụ để kiểm tra liệu việc cân bằng cơ hội có tạo ra chênh lệch về giá trị cuốc hay không.")

    add_heading(doc, "2.6. Khoảng trống nghiên cứu", 2)
    add_body(doc, "Các công trình trước cung cấp nền tảng về định tuyến có cửa sổ thời gian, ghép nối động, điều phối theo Batch, tối ưu đa mục tiêu và các cách tiếp cận công bằng khác nhau [1]–[5]. Tuy nhiên, trong phạm vi các tài liệu được khảo sát, vẫn cần đánh giá một quy trình kết hợp đồng thời cơ chế Batching, mô hình ghép một tài xế–một yêu cầu có cửa sổ thời gian, mục tiêu quãng đường chạy rỗng và công bằng theo cơ hội nhận cuốc trong một khung NSGA-II thống nhất.")
    add_body(doc, "Từ khoảng trống đó, đề tài thiết kế mã hóa số nguyên, cơ chế kiểm tra–sửa nghiệm và mô phỏng nhiều Batch liên tiếp để tạo tập nghiệm Pareto giữa hiệu quả vận hành và công bằng. Vấn đề cần kiểm chứng không phải là NSGA-II luôn vượt trội, mà là thuật toán có tạo được tập nghiệm đánh đổi hữu ích, duy trì nghiệm khả thi và hoàn thành trong giới hạn thời gian của một Batch hay không.")

    add_heading(doc, "3. MỤC ĐÍCH VÀ NHIỆM VỤ NGHIÊN CỨU", 1)
    add_heading(doc, "3.1. Mục đích nghiên cứu", 2)
    add_body(doc, "Mục đích của đề tài là xây dựng và đánh giá một phương pháp điều phối gọi xe đa mục tiêu dựa trên NSGA-II trong cơ chế Batching có ràng buộc cửa sổ thời gian. Phương pháp dự kiến tạo ra tập nghiệm Pareto thể hiện sự đánh đổi giữa tổng quãng đường chạy rỗng và mức chênh lệch cơ hội nhận cuốc giữa các tài xế, đồng thời bảo đảm các điều kiện ghép nối trong mỗi Batch.")
    add_heading(doc, "3.2. Nhiệm vụ nghiên cứu", 2)
    for item in [
        "Hệ thống hóa cơ sở lý thuyết về bài toán điều phối gọi xe, cơ chế Batching, tối ưu đa mục tiêu, biên Pareto và thuật toán NSGA-II.",
        "Xác định tập dữ liệu, tham số, biến quyết định, hai hàm mục tiêu và các ràng buộc của mô hình ghép nối một tài xế–một yêu cầu.",
        "Thiết kế mã hóa nhiễm sắc thể, khởi tạo quần thể, lựa chọn, lai ghép, đột biến và cơ chế kiểm tra–sửa nghiệm cho NSGA-II.",
        "Xây dựng môi trường mô phỏng bằng Python, tổ chức dữ liệu theo các Batch liên tiếp và cập nhật trạng thái tài xế sau mỗi quyết định.",
        "Thực hiện các kịch bản giờ cao điểm, thấp điểm và phân tích độ nhạy theo độ dài Batch; so sánh với phương pháp Greedy tài xế gần nhất.",
        "Đánh giá kết quả theo quãng đường chạy rỗng, thời gian chờ, tỷ lệ phục vụ, độ lệch chuẩn tỷ lệ tiếp cận cuốc, phân phối thu nhập, thời gian tính toán và chất lượng tập nghiệm Pareto."
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4. ĐỐI TƯỢNG VÀ PHẠM VI NGHIÊN CỨU", 1)
    add_heading(doc, "4.1. Đối tượng nghiên cứu", 2)
    add_body(doc, "Đối tượng nghiên cứu là bài toán tối ưu đa mục tiêu trong điều phối phương tiện cho dịch vụ gọi xe công nghệ theo cơ chế Batching; thuật toán NSGA-II và các thành phần được điều chỉnh cho bài toán ghép nối; cùng khung mô phỏng dùng để đánh giá sự cân bằng giữa hiệu quả vận hành và công bằng giữa các tài xế.")
    add_heading(doc, "4.2. Phạm vi nội dung", 2)
    add_body(doc, "Mô hình cơ sở xét một tài xế phục vụ tối đa một yêu cầu và một yêu cầu được gán tối đa cho một tài xế trong mỗi Batch. Mục tiêu thứ nhất là tối thiểu hóa tổng quãng đường chạy rỗng từ vị trí tài xế đến điểm đón. Mục tiêu thứ hai là tối thiểu hóa độ lệch chuẩn của tỷ lệ tiếp cận cuốc giữa các tài xế. Tỷ lệ này được tính bằng số cuốc đã nhận trên số lần tài xế thực sự đủ điều kiện nhận ít nhất một cuốc; thu nhập được ghi nhận như chỉ số đánh giá bổ trợ, không tham gia hàm mục tiêu.")
    add_body(doc, "Ràng buộc chính gồm tính duy nhất của phép gán và giới hạn thời gian đến điểm đón. Vì mô hình không xét đi chung xe hoặc một tài xế phục vụ chuỗi nhiều yêu cầu trong cùng Batch, ràng buộc sức chứa và thứ tự đón–trả không được đưa vào mô hình cơ sở; đây là hướng mở rộng sau nghiên cứu.")
    add_heading(doc, "4.3. Phạm vi dữ liệu và kịch bản", 2)
    add_body(doc, "Nghiên cứu dự kiến sử dụng một bộ dữ liệu hành trình thực tế có các trường thời gian, tọa độ điểm đón và điểm trả. Trong hai lựa chọn được nêu ở tài liệu nguồn là NYC Taxi và RideAustin, bộ dữ liệu chính thức sẽ được chốt sau khi kiểm tra mức độ đầy đủ của thời điểm yêu cầu và tính phù hợp với mô hình. Phạm vi thử nghiệm dự kiến giới hạn ở một khu vực đô thị và một khoảng thời gian đại diện, có phân tách giờ cao điểm và thấp điểm.")
    add_heading(doc, "4.4. Giả định và giới hạn", 2)
    for item in [
        "Thời gian di chuyển được ước lượng từ dữ liệu lịch sử hoặc từ khoảng cách và hệ số vận tốc theo kịch bản; không mô phỏng sự cố giao thông đột ngột.",
        "Hành khách không hủy chuyến sau khi đã được gán; trạng thái tài xế được cập nhật chính xác sau mỗi Batch.",
        "Độ dài Batch mặc định là 60 giây; giá trị 120 giây chỉ dùng cho phân tích độ nhạy, không phải cấu hình mặc định cạnh tranh.",
        "Kết quả mô phỏng chỉ cho phép kết luận trong phạm vi dữ liệu, kịch bản và tham số đã thử nghiệm."
    ]:
        add_bullet(doc, item)

    add_heading(doc, "5. PHƯƠNG PHÁP NGHIÊN CỨU", 1)
    add_heading(doc, "5.1. Thiết kế nghiên cứu", 2)
    add_body(doc, "Nghiên cứu sử dụng phương pháp mô hình hóa toán học kết hợp thiết kế thuật toán và thực nghiệm mô phỏng. Quy trình gồm sáu bước: khảo sát tài liệu; chuẩn hóa dữ liệu; xây dựng mô hình; cài đặt NSGA-II và phương pháp đối chứng; chạy mô phỏng theo Batch; phân tích kết quả và kiểm định các giả thuyết.")
    add_heading(doc, "5.2. Cơ chế Batching", 2)
    add_body(doc, "Dòng yêu cầu được chia thành các khoảng có độ dài ΔT. Tại cuối Batch thứ k, hệ thống thu thập tập yêu cầu chưa được phục vụ R và tập tài xế đang sẵn sàng D, cố định trạng thái của chúng trong thời gian giải bài toán, chọn một nghiệm điều phối rồi cập nhật trạng thái trước Batch kế tiếp. Cấu hình mặc định là ΔT = 60 giây; ΔT = 120 giây được dùng để phân tích ảnh hưởng của thời gian gom yêu cầu.")
    add_heading(doc, "5.3. Mô hình toán học đa mục tiêu", 2)
    add_heading(doc, "5.3.1. Tập hợp và biến quyết định", 3)
    add_body(doc, "Gọi R là tập yêu cầu trong Batch hiện tại và D là tập tài xế sẵn sàng. Với mỗi cặp (i, j), biến nhị phân xᵢⱼ nhận giá trị 1 nếu tài xế i được gán cho yêu cầu j và nhận giá trị 0 trong trường hợp còn lại. Ký hiệu dist(dᵢ, pⱼ) là khoảng cách từ vị trí hiện tại của tài xế i đến điểm đón pⱼ.")
    add_heading(doc, "5.3.2. Hàm mục tiêu", 3)
    add_body(doc, "Hàm mục tiêu thứ nhất tối thiểu hóa tổng quãng đường chạy rỗng của các tài xế được phân công:")
    add_equation(doc, "eq_5_1.png", "5.1", 9.5)
    add_body(doc, "Để biểu diễn công bằng theo cơ hội nhận cuốc, gọi eᵢ là số Batch tính đến thời điểm hiện tại mà tài xế i đang sẵn sàng và có ít nhất một yêu cầu khả thi; nᵢ là số cuốc tích lũy mà tài xế i đã được phân sau khi áp dụng phương án X. Gọi D⁺ là tập tài xế có eᵢ > 0, qᵢ là tỷ lệ tiếp cận cuốc của tài xế i và q̄ là tỷ lệ trung bình:")
    add_equation(doc, "eq_5_2.png", "5.2", 5.6)
    add_body(doc, "Hàm mục tiêu thứ hai tối thiểu hóa độ lệch chuẩn của tỷ lệ tiếp cận cuốc giữa các tài xế đã có ít nhất một cơ hội khả thi:")
    add_equation(doc, "eq_5_3.png", "5.3", 7.8)
    add_body(doc, "Bài toán đa mục tiêu được biểu diễn dưới dạng:")
    add_equation(doc, "eq_5_4.png", "5.4", 4.6)
    add_body(doc, "Hai mục tiêu được giữ độc lập trong quá trình xếp hạng Pareto. Nghiên cứu không dùng tổng trọng số làm cơ chế tối ưu chính; nếu cần chọn một nghiệm để vận hành, quy tắc chọn nghiệm sẽ được công bố trước khi đánh giá, chẳng hạn chọn điểm gối hoặc áp dụng ngưỡng dịch vụ tối thiểu.")
    add_heading(doc, "5.3.3. Các ràng buộc", 3)
    add_body(doc, "Mỗi yêu cầu được gán cho tối đa một tài xế:")
    add_equation(doc, "eq_5_5.png", "5.5", 5.3)
    add_body(doc, "Mỗi tài xế nhận tối đa một yêu cầu trong Batch:")
    add_equation(doc, "eq_5_6.png", "5.6", 5.3)
    add_body(doc, "Nếu tài xế i được gán cho yêu cầu j, thời điểm đến điểm đón không được vượt quá giới hạn chờ tối đa của yêu cầu:")
    add_equation(doc, "eq_5_7.png", "5.7", 8.8)
    add_body(doc, "Yêu cầu không tìm được tài xế khả thi có thể được giữ lại cho Batch kế tiếp hoặc ghi nhận là không được phục vụ theo quy tắc thí nghiệm. Quy tắc này phải được áp dụng giống nhau cho NSGA-II và Greedy để bảo đảm so sánh công bằng.")
    add_heading(doc, "5.4. Thiết kế NSGA-II", 2)
    add_heading(doc, "5.4.1. Mã hóa và khởi tạo", 3)
    add_body(doc, "Một nhiễm sắc thể có độ dài bằng số yêu cầu trong Batch; gen thứ j chứa định danh tài xế được gán cho yêu cầu j. Giá trị 0 biểu diễn yêu cầu chưa được phục vụ. Quần thể ban đầu dự kiến kết hợp nghiệm ngẫu nhiên khả thi với một số nghiệm được gieo từ heuristic khoảng cách để giảm tỷ lệ nghiệm kém chất lượng mà không làm mất tính đa dạng.")
    add_heading(doc, "5.4.2. Xếp hạng không trội và khoảng cách mật độ", 3)
    add_body(doc, "Với hai nghiệm X⁽ᵃ⁾ và X⁽ᵇ⁾, X⁽ᵃ⁾ trội hơn X⁽ᵇ⁾ khi không kém ở mọi mục tiêu và tốt hơn ở ít nhất một mục tiêu:")
    add_equation(doc, "eq_5_8.png", "5.8", 12.0)
    add_body(doc, "Các nghiệm được phân thành các tầng Pareto bằng sắp xếp không trội nhanh. Trong cùng một tầng, khoảng cách mật độ được dùng để ưu tiên các nghiệm ở vùng thưa và duy trì sự đa dạng:")
    add_equation(doc, "eq_5_9.png", "5.9", 7.7)
    add_heading(doc, "5.4.3. Toán tử tiến hóa và sửa nghiệm", 3)
    add_body(doc, "Lựa chọn giải đấu nhị phân ưu tiên nghiệm có hạng Pareto tốt hơn; nếu cùng hạng, nghiệm có khoảng cách mật độ lớn hơn được chọn. Lai ghép hai điểm và đột biến hoán vị được áp dụng trên mã hóa số nguyên. Sau mỗi toán tử, cơ chế sửa nghiệm lần lượt loại gán trùng tài xế, loại cặp vi phạm cửa sổ thời gian và thử tái phân công các yêu cầu có giá trị 0 bằng tài xế còn rảnh và khả thi. Các tham số kích thước quần thể, xác suất lai ghép, xác suất đột biến và số thế hệ sẽ được hiệu chỉnh trên tập phát triển, tách khỏi tập đánh giá cuối.")
    add_heading(doc, "5.5. Thực nghiệm và đánh giá", 2)
    add_body(doc, "Dữ liệu sẽ được làm sạch, sắp xếp theo thời gian và chia thành các luồng yêu cầu theo kịch bản giờ cao điểm và thấp điểm. Phương pháp đối chứng là Greedy chọn tài xế khả thi gần điểm đón nhất. Mỗi cấu hình ngẫu nhiên của NSGA-II cần được chạy lặp lại với nhiều hạt giống; báo cáo giá trị trung bình, độ phân tán và khoảng tin cậy khi phù hợp.")
    add_body(doc, "Các chỉ số đánh giá gồm tổng quãng đường chạy rỗng, thời gian chờ trung bình, tỷ lệ yêu cầu được phục vụ, độ lệch chuẩn tỷ lệ tiếp cận cuốc, phân phối số cuốc, thời gian tính toán mỗi Batch và tỷ lệ nghiệm khả thi. Thu nhập tích lũy và độ lệch thu nhập được báo cáo như chỉ số phụ để kiểm tra liệu việc cân bằng cơ hội có dẫn đến chênh lệch giá trị cuốc hay không. Chất lượng tập nghiệm đa mục tiêu dự kiến được mô tả bằng biểu đồ Pareto và ít nhất một chỉ số chất lượng phù hợp sau khi xác định điểm tham chiếu. Phân tích độ nhạy sẽ xem xét ΔT = 60 và 120 giây cùng một số quy mô Batch đại diện.")

    add_heading(doc, "6. GIẢ THUYẾT KHOA HỌC", 1)
    add_heading(doc, "6.1. Giả thuyết về hiệu quả của Batching", 2)
    add_body(doc, "H₁: Trong cùng dữ liệu và điều kiện mô phỏng, cơ chế ghép nối theo Batch kết hợp NSGA-II dự kiến tạo ra ít nhất một phương án có tổng quãng đường chạy rỗng thấp hơn hoặc độ lệch tỷ lệ tiếp cận cuốc nhỏ hơn Greedy, trong khi vẫn duy trì tỷ lệ phục vụ và cửa sổ thời gian ở mức chấp nhận được. Giả thuyết được kiểm tra bằng so sánh các chỉ số trên nhiều lần chạy.")
    add_heading(doc, "6.2. Giả thuyết về tập nghiệm Pareto", 2)
    add_body(doc, "H₂: NSGA-II dự kiến tạo được tập nghiệm không trội có sự phân bố đủ để quan sát đánh đổi giữa F₁ và F₂. Giả thuyết được kiểm tra qua biên Pareto, độ phân tán nghiệm và chỉ số chất lượng đa mục tiêu được chọn.")
    add_heading(doc, "6.3. Giả thuyết về cơ chế sửa nghiệm", 2)
    add_body(doc, "H₃: Cơ chế sửa nghiệm dự kiến làm tăng tỷ lệ cá thể khả thi so với cấu hình chỉ dùng hàm phạt, đồng thời không làm suy giảm nghiêm trọng tính đa dạng của quần thể. Giả thuyết được kiểm tra bằng thí nghiệm loại bỏ thành phần, so sánh tỷ lệ nghiệm khả thi, thời gian hội tụ và chất lượng tập Pareto.")
    add_heading(doc, "6.4. Giả thuyết về khả năng đáp ứng thời gian", 2)
    add_body(doc, "H₄: Với quy mô Batch và tham số được lựa chọn trong phạm vi đề tài, thời gian tính toán trung bình của thuật toán dự kiến nhỏ hơn độ dài Batch:")
    add_equation(doc, "eq_6_1.png", "6.1", 3.7)
    add_body(doc, "Giả thuyết được chấp nhận trong phạm vi thử nghiệm nếu thời gian tính toán được báo cáo theo phân phối và tỷ lệ Batch đáp ứng giới hạn, không chỉ dựa trên một lần chạy.")

    add_heading(doc, "7. NHỮNG ĐÓNG GÓP MỚI DỰ KIẾN CỦA ĐỀ TÀI", 1)
    add_body(doc, "Các đóng góp dưới đây là đóng góp dự kiến và chỉ được xác nhận sau khi nghiên cứu hoàn tất:")
    for item in [
        "Một mô hình ghép nối gọi xe theo Batch tích hợp đồng thời mục tiêu quãng đường chạy rỗng và độ lệch chuẩn tỷ lệ tiếp cận cuốc, có ràng buộc cửa sổ thời gian.",
        "Một cách mã hóa số nguyên và cơ chế sửa nghiệm dành cho cấu trúc ghép một tài xế–một yêu cầu, giúp kiểm soát gán trùng và cặp ghép không khả thi.",
        "Một quy trình mô phỏng nhiều Batch liên tiếp, có cập nhật trạng thái tài xế và ghi nhận đồng thời các chỉ số vận hành, dịch vụ, công bằng và thời gian tính toán.",
        "Một đánh giá thực nghiệm có đối chứng với Greedy và phân tích độ nhạy theo độ dài Batch, qua đó làm rõ điều kiện mà NSGA-II có thể tạo ra phương án đánh đổi hữu ích."
    ]:
        add_bullet(doc, item)

    add_heading(doc, "8. DỰ KIẾN KẾ HOẠCH NGHIÊN CỨU", 1)
    add_body(doc, "Kế hoạch được xây dựng theo sáu tháng tương đối vì tài liệu nguồn chưa cung cấp mốc lịch học chính thức. Khi có lịch của khoa, các mốc “Tháng 1–6” cần được thay bằng tháng hoặc tuần cụ thể.")
    caption = doc.add_paragraph()
    format_para(caption, align=WD_ALIGN_PARAGRAPH.CENTER, first=None, before=4, after=4, keep=True)
    set_run_font(caption.add_run("Bảng 8.1. Kế hoạch nghiên cứu dự kiến"), 13, True)
    rows = [
        ("1", "Tổng quan tài liệu; xác định câu hỏi, khoảng trống và tiêu chí đánh giá.", "Tháng 1", "Khung lý thuyết; danh mục nguồn; đề cương chi tiết."),
        ("2", "Chốt bộ dữ liệu; làm sạch dữ liệu; hoàn thiện mô hình toán học và quy tắc Batch.", "Tháng 2", "Tập dữ liệu sạch; đặc tả mô hình, biến và ràng buộc."),
        ("3", "Cài đặt Greedy, môi trường mô phỏng và các thành phần NSGA-II.", "Tháng 3", "Mã nguồn chạy được; bộ kiểm thử ràng buộc."),
        ("4", "Hoàn thiện toán tử, cơ chế sửa nghiệm; hiệu chỉnh tham số trên tập phát triển.", "Tháng 4", "Cấu hình thực nghiệm; báo cáo kiểm thử và hiệu chỉnh."),
        ("5", "Chạy thí nghiệm cao điểm, thấp điểm và độ nhạy; tổng hợp số liệu.", "Tháng 5", "Tập kết quả; biểu đồ Pareto; bảng so sánh với Greedy."),
        ("6", "Phân tích kết quả; hoàn thiện báo cáo; rà soát trích dẫn và chuẩn bị bảo vệ.", "Tháng 6", "Báo cáo ĐACN; mã nguồn và phụ lục; slide bảo vệ.")
    ]
    table = add_table(doc, ["Giai đoạn", "Nội dung công việc", "Thời gian", "Sản phẩm dự kiến"], rows, [1.6, 5.5, 2.2, 4.2])
    set_table_borders(table)
    add_body(doc, "Rủi ro chính gồm dữ liệu không có trường thời điểm yêu cầu phù hợp, quy mô dữ liệu quá lớn và thời gian tiến hóa vượt giới hạn Batch. Biện pháp dự phòng là chọn bộ dữ liệu có cấu trúc thích hợp, giới hạn phạm vi theo khu vực và thời gian, tiền tính toán ma trận khoảng cách khi hợp lý, đồng thời hiệu chỉnh kích thước quần thể và số thế hệ trên tập phát triển. Mọi tối ưu hóa mã nguồn phải giữ nguyên logic so sánh giữa NSGA-II và phương pháp đối chứng.")

    add_heading(doc, "9. DỰ KIẾN NỘI DUNG CỦA ĐỒ ÁN CHUYÊN NGÀNH", 1)
    add_body(doc, "Tệp nguồn dành cho phần này chưa có nội dung. Cấu trúc dưới đây được xây dựng từ mục tiêu, mô hình và kế hoạch nghiên cứu đã nêu, gồm ba chương theo mẫu hướng dẫn:")
    add_heading(doc, "Chương 1. Cơ sở lý thuyết và tổng quan nghiên cứu", 2)
    for item in [
        "1.1. Bài toán điều phối gọi xe và các yêu cầu vận hành.",
        "1.2. Cơ chế Batching và ràng buộc cửa sổ thời gian.",
        "1.3. Tối ưu đa mục tiêu, quan hệ trội và biên Pareto.",
        "1.4. Thuật toán NSGA-II và các nghiên cứu liên quan.",
        "1.5. Khoảng trống nghiên cứu và định hướng của đề tài."
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Chương 2. Mô hình bài toán và phương pháp đề xuất", 2)
    for item in [
        "2.1. Phát biểu bài toán, giả định, tập hợp và biến quyết định.",
        "2.2. Hai hàm mục tiêu và các ràng buộc.",
        "2.3. Mô hình vận hành theo Batch.",
        "2.4. Mã hóa, toán tử tiến hóa và cơ chế sửa nghiệm.",
        "2.5. Quy trình NSGA-II và quy tắc chọn nghiệm điều phối.",
        "2.6. Thiết kế môi trường mô phỏng và phương pháp đối chứng."
    ]:
        add_bullet(doc, item)
    add_heading(doc, "Chương 3. Thực nghiệm, kết quả và thảo luận", 2)
    for item in [
        "3.1. Bộ dữ liệu, tiền xử lý và thiết lập kịch bản.",
        "3.2. Cấu hình thuật toán và quy trình lặp thực nghiệm.",
        "3.3. Kết quả về hiệu quả vận hành và chất lượng dịch vụ.",
        "3.4. Kết quả về công bằng và tập nghiệm Pareto.",
        "3.5. So sánh với Greedy, phân tích độ nhạy và kiểm định giả thuyết.",
        "3.6. Hạn chế, kết luận và hướng phát triển."
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10. DANH MỤC TÀI LIỆU THAM KHẢO", 1)
    add_reference_placeholder(doc, "[1] G. B. Dantzig and J. H. Ramser, “The truck dispatching problem,” Management Science, vol. 6, no. 1, pp. 80–91, 1959, doi: 10.1287/mnsc.6.1.80.")
    add_reference_placeholder(doc, "[2] M. M. Solomon, “Algorithms for the vehicle routing and scheduling problems with time window constraints,” Operations Research, vol. 35, no. 2, pp. 254–265, 1987, doi: 10.1287/opre.35.2.254.")
    add_reference_placeholder(doc, "[3] K. Deb, A. Pratap, S. Agarwal, and T. Meyarivan, “A fast and elitist multiobjective genetic algorithm: NSGA-II,” IEEE Transactions on Evolutionary Computation, vol. 6, no. 2, pp. 182–197, Apr. 2002, doi: 10.1109/4235.996017.")
    add_reference_placeholder(doc, "[4] Z. Qin et al., “Ride-hailing order dispatching at DiDi via reinforcement learning,” INFORMS Journal on Applied Analytics, vol. 50, no. 5, pp. 272–286, 2020, doi: 10.1287/inte.2020.1047.")
    add_reference_placeholder(doc, "[5] S. Ben-Gal and M. Tzur, “Data-driven policies for the online ride-hailing problem with fairness,” Transportation Science, vol. 59, no. 3, pp. 647–669, 2025, doi: 10.1287/trsc.2023.0068.")
    add_reference_placeholder(doc, "[6] [Cần bổ sung tài liệu mô tả phương pháp Greedy hoặc baseline được sử dụng trong thực nghiệm.]")
    add_reference_placeholder(doc, "[7] [Cần bổ sung trang tài liệu chính thức và phiên bản của bộ dữ liệu được chọn: NYC Taxi hoặc RideAustin.]")

    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)

    core = doc.core_properties
    core.title = "Đề cương nghiên cứu NSGA-II cho điều phối gọi xe đa mục tiêu"
    core.subject = "Đề cương đồ án chuyên ngành"
    core.keywords = "NSGA-II, ride-hailing, batching, time windows, opportunity fairness, multi-objective optimization"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
