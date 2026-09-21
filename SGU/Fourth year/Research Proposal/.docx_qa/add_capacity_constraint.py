from copy import deepcopy
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "De_cuong_nghien_cuu_NSGAII_Ride_Hailing.docx"
OUTPUT = Path(__import__("os").environ["DOCX_OUT"])


def find_one(document, prefix):
    matches = [p for p in document.paragraphs if p.text.startswith(prefix)]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def replace_single_run(paragraph, text):
    if len(paragraph.runs) != 1:
        raise RuntimeError(f"Expected a single-run paragraph, found {len(paragraph.runs)} runs")
    paragraph.runs[0].text = text


def insert_matching_paragraph_before(reference, text):
    new_paragraph = reference.insert_paragraph_before()
    new_paragraph.style = reference.style
    if reference._p.pPr is not None:
        if new_paragraph._p.pPr is not None:
            new_paragraph._p.remove(new_paragraph._p.pPr)
        new_paragraph._p.insert(0, deepcopy(reference._p.pPr))
    run = new_paragraph.add_run(text)
    if reference.runs and reference.runs[0]._r.rPr is not None:
        run._r.insert(0, deepcopy(reference.runs[0]._r.rPr))
    return new_paragraph


doc = Document(SOURCE)

scope = find_one(doc, "Ràng buộc chính gồm tính duy nhất của phép gán")
replace_single_run(
    scope,
    "Ràng buộc chính gồm tính duy nhất của phép gán, giới hạn sức chứa và giới hạn thời gian đến điểm đón. "
    "Vì mô hình không xét đi chung xe hoặc một tài xế phục vụ chuỗi nhiều yêu cầu trong cùng Batch, "
    "các ràng buộc tải trọng biến đổi và sắp xếp nhiều điểm đón–trả không được đưa vào mô hình cơ sở; "
    "đây là hướng mở rộng sau nghiên cứu.",
)

variables = find_one(doc, "Gọi R là tập yêu cầu trong Batch hiện tại")
replace_single_run(
    variables,
    variables.text
    + " Ký hiệu sⱼ là số hành khách của yêu cầu j và Qᵢ là sức chứa tối đa của phương tiện thuộc tài xế i; "
      "xe máy có Qᵢ = 1, còn ô tô có Qᵢ bằng số chỗ hành khách được phép chở.",
)

pickup_time = find_one(doc, "Nếu tài xế i được gán cho yêu cầu j, thời điểm đến điểm đón")
insert_matching_paragraph_before(
    pickup_time,
    "Ràng buộc sức chứa được áp dụng cho từng cặp ghép: yêu cầu j chỉ có thể được gán cho tài xế i khi "
    "sⱼ ≤ Qᵢ. Tương đương, xᵢⱼ = 0 đối với mọi cặp (i, j) có sⱼ > Qᵢ. Vì vậy, yêu cầu có nhiều hơn "
    "một hành khách không thể được gán cho xe máy, nhưng có thể được gán cho ô tô nếu không vượt quá sức chứa của xe.",
)

doc.save(OUTPUT)
print(OUTPUT)
